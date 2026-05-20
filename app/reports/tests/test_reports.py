"""
test_reports.py
===============
ครอบคลุม:
  - ReportFilterParams validation
  - summary statistics aggregation
  - CSV / Excel export bytes
  - API endpoints (permission + response)
"""

import csv
import io
from datetime import date, timedelta
from unittest.mock import patch

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from rest_framework.exceptions import ValidationError
from rest_framework.test import APIClient

import openpyxl

from account.models import User
from bookings.models import Booking, TeachingInfo, TrainingInfo
from reports.filters import ReportFilterParams
from reports.services.export_service import build_csv_bytes, build_excel_bytes
from rooms.models import Room


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def make_user(username: str, role: str = "lecturer") -> User:
    user = User(username=username, role=role)
    user.set_unusable_password()
    user.save()
    return user


def make_room(code: str = "ECE-101") -> Room:
    return Room.objects.create(
        room_code=code,
        room_name=f"ห้อง {code}",
        room_type="Classroom",
        capacity=30,
    )


def make_booking(
    room: Room,
    booker: User,
    purpose: str = "teaching",
    status: str = "Approved",
    days_offset: int = 0,
    program_type: str = "Bachelor",
) -> Booking:
    now = timezone.now() + timedelta(days=days_offset)
    bk = Booking.objects.create(
        room=room,
        booker=booker,
        start_datetime=now,
        end_datetime=now + timedelta(hours=2),
        status=status,
        purpose_type=purpose,
    )
    if purpose == "teaching":
        TeachingInfo.objects.create(
            booking=bk,
            subject_code="CN334",
            subject_name="Software Engineering",
            program_type=program_type,
        )
    elif purpose == "training":
        TrainingInfo.objects.create(booking=bk, topic="Python Workshop")
    return bk


# ---------------------------------------------------------------------------
# Filter validation tests
# ---------------------------------------------------------------------------

class TestReportFilterParams(TestCase):

    def test_valid_params(self):
        f = ReportFilterParams.from_query_params(
            {"date_from": "2024-01-01", "date_to": "2024-03-31", "purpose_type": "teaching"}
        )
        self.assertEqual(f.date_from, date(2024, 1, 1))
        self.assertEqual(f.purpose_type, "teaching")

    def test_invalid_date_format(self):
        with self.assertRaises(ValidationError) as ctx:
            ReportFilterParams.from_query_params({"date_from": "01-01-2024"})
        self.assertIn("date_from", ctx.exception.detail)

    def test_date_range_reversed(self):
        with self.assertRaises(ValidationError) as ctx:
            ReportFilterParams.from_query_params(
                {"date_from": "2024-12-31", "date_to": "2024-01-01"}
            )
        self.assertIn("date_range", ctx.exception.detail)

    def test_invalid_purpose_type(self):
        with self.assertRaises(ValidationError) as ctx:
            ReportFilterParams.from_query_params({"purpose_type": "lecture"})
        self.assertIn("purpose_type", ctx.exception.detail)

    def test_program_type_without_teaching_raises(self):
        with self.assertRaises(ValidationError) as ctx:
            ReportFilterParams.from_query_params(
                {"purpose_type": "training", "program_type": "Bachelor"}
            )
        self.assertIn("program_type", ctx.exception.detail)

    def test_empty_params_returns_defaults(self):
        f = ReportFilterParams.from_query_params({})
        self.assertIsNone(f.date_from)
        self.assertIsNone(f.purpose_type)


# ---------------------------------------------------------------------------
# Summary statistics tests
# ---------------------------------------------------------------------------

class TestSummaryStatistics(TestCase):

    def setUp(self):
        self.room = make_room()
        self.lecturer = make_user("lecturer1")
        make_booking(self.room, self.lecturer, purpose="teaching", program_type="Bachelor")
        make_booking(self.room, self.lecturer, purpose="teaching", program_type="Master")
        make_booking(self.room, self.lecturer, purpose="training")

    def test_total_count(self):
        from reports.services.report_query_service import get_summary_statistics
        f = ReportFilterParams()
        stats = get_summary_statistics(f)
        self.assertEqual(stats["total_bookings"], 3)

    def test_by_purpose(self):
        from reports.services.report_query_service import get_summary_statistics
        f = ReportFilterParams()
        stats = get_summary_statistics(f)
        self.assertEqual(stats["by_purpose"].get("teaching"), 2)
        self.assertEqual(stats["by_purpose"].get("training"), 1)

    def test_by_program_type(self):
        from reports.services.report_query_service import get_summary_statistics
        f = ReportFilterParams()
        stats = get_summary_statistics(f)
        self.assertEqual(stats["by_program_type"].get("Bachelor"), 1)
        self.assertEqual(stats["by_program_type"].get("Master"), 1)

    def test_filter_by_purpose(self):
        from reports.services.report_query_service import get_summary_statistics
        f = ReportFilterParams(purpose_type="training")
        stats = get_summary_statistics(f)
        self.assertEqual(stats["total_bookings"], 1)

    def test_monthly_trend_sorted(self):
        from reports.services.report_query_service import get_summary_statistics
        f = ReportFilterParams()
        stats = get_summary_statistics(f)
        months = [t["month"] for t in stats["monthly_trend"]]
        self.assertEqual(months, sorted(months))


# ---------------------------------------------------------------------------
# Export service tests
# ---------------------------------------------------------------------------

SAMPLE_HEADERS = [
    "รหัสการจอง", "รหัสห้อง", "ชื่อห้อง", "Username ผู้จอง", "ชื่อผู้จอง", "ภาควิชา",
    "วันเวลาเริ่มต้น", "วันเวลาสิ้นสุด", "ชั่วโมงการใช้งาน", "สถานะ", "ประเภทการใช้งาน",
    "รหัสวิชา", "ชื่อวิชา", "หลักสูตร", "หัวข้ออบรม", "คำขอเพิ่มเติม", "วันที่จอง"
]

SAMPLE_ROWS = [
    [
        1, "ECE-101", "ห้องเรียน 101", "user1", "นาย ทดสอบ", "วิศวกรรมไฟฟ้า",
        "2024-01-15 09:00", "2024-01-15 11:00", 2.0, "อนุมัติแล้ว", "สอน",
        "CN334", "Software Engineering", "ปริญญาตรีภาคปกติ", "-", "", "2024-01-10 08:00"
    ]
]


class TestExportService(TestCase):

    def test_csv_bytes_not_empty(self):
        result = build_csv_bytes(SAMPLE_HEADERS, SAMPLE_ROWS)
        self.assertIsInstance(result, bytes)
        self.assertGreater(len(result), 0)

    def test_csv_has_thai_header(self):
        result = build_csv_bytes(SAMPLE_HEADERS, SAMPLE_ROWS)
        text = result.decode("utf-8-sig")
        self.assertIn("รหัสการจอง", text)

    def test_csv_empty_rows_returns_header_only(self):
        result = build_csv_bytes(SAMPLE_HEADERS, [])
        text = result.decode("utf-8-sig")
        lines = [l for l in text.strip().splitlines() if l]
        self.assertEqual(len(lines), 1)  # header only

    def test_excel_bytes_valid_workbook(self):
        result = build_excel_bytes(SAMPLE_HEADERS, SAMPLE_ROWS)
        wb = openpyxl.load_workbook(io.BytesIO(result))
        ws = wb.active
        # row 1 = header, row 2 = data
        self.assertEqual(ws.max_row, 2)

    def test_excel_header_bold(self):
        result = build_excel_bytes(SAMPLE_HEADERS, SAMPLE_ROWS)
        wb = openpyxl.load_workbook(io.BytesIO(result))
        ws = wb.active
        self.assertTrue(ws.cell(row=1, column=1).font.bold)


# ---------------------------------------------------------------------------
# API endpoint tests
# ---------------------------------------------------------------------------

class TestReportAPIPermissions(TestCase):

    def setUp(self):
        self.client = APIClient()
        self.admin = make_user("admin1", role="admin")
        self.lecturer = make_user("lec1", role="lecturer")

    def _login(self, user: User):
        self.client.force_login(user)

    def test_summary_requires_login(self):
        response = self.client.get(reverse("reports:summary"))
        self.assertEqual(response.status_code, 403)

    def test_summary_rejects_non_admin(self):
        self._login(self.lecturer)
        response = self.client.get(reverse("reports:summary"))
        self.assertEqual(response.status_code, 403)

    def test_summary_allows_admin(self):
        self._login(self.admin)
        response = self.client.get(reverse("reports:summary"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("statistics", response.data)

    def test_export_csv_content_type(self):
        self._login(self.admin)
        response = self.client.get(reverse("reports:export_csv"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("text/csv", response["Content-Type"])

    def test_export_excel_content_type(self):
        self._login(self.admin)
        response = self.client.get(reverse("reports:export_excel"))
        self.assertEqual(response.status_code, 200)
        self.assertIn(
            "spreadsheetml",
            response["Content-Type"],
        )

    def test_summary_invalid_filter_returns_400(self):
        self._login(self.admin)
        response = self.client.get(
            reverse("reports:summary"), {"date_from": "not-a-date"}
        )
        self.assertEqual(response.status_code, 400)